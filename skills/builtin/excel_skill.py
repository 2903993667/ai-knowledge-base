"""Excel (.xlsx) parser skill."""
import openpyxl

SKILL = {
    "name": "Excel",
    "description": "Parse Excel .xlsx files into records",
    "extensions": [".xlsx"],
    "mime_types": ["application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"],
}

async def parse(file_path: str = "", data: bytes = None):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    all_records = []
    fields = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
        if not fields:
            fields = headers
        for row in rows[1:]:
            record = {"_sheet": sheet_name}
            for h, v in zip(headers, row):
                record[h] = str(v) if v is not None else ""
            all_records.append(record)

    wb.close()
    return {"records": all_records, "fields": fields, "metadata": {"row_count": len(all_records), "format": "xlsx", "sheets": wb.sheetnames if hasattr(wb, 'sheetnames') else []}}
